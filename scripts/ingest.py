#!/usr/bin/env python3
"""
BIS Mitra — data ingestion.

Reads the supplied seed workbooks and emits a single normalized knowledge layer
at  data/knowledge.json  that the application consumes.

    Excel workbooks
        -> data loader (this script)
        -> normalized objects
        -> repository (src/lib/repository.ts)
        -> retrieval engine
        -> evidence packet
        -> UI / LLM

Trust rules enforced here (mirrors the workbook's own trust policy):
  * No standard, QCO date, licence number, clause, page or source URL is invented.
  * Missing values stay null. Nothing is back-filled from memory.
  * Every field keeps its provenance (workbook + sheet it came from).

Usage:  python3 scripts/ingest.py
"""

from __future__ import annotations

import json
import re
import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATASETS = ROOT / "datasets"
OUT = ROOT / "data" / "knowledge.json"

PRIMARY = "BIS_Navigator_Seed_Data_Normalized(2).xlsx"
CONSUMER = "BIS_Navigator_Seed_Data_With_Consumer.xlsx"
REFERENCE = "Total Published Standards (1).xlsx"

report: list[str] = []
problems: list[str] = []


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def load(name: str) -> Any:
    path = DATASETS / name
    if not path.exists():
        problems.append(f"missing workbook: {name}")
        return None
    return openpyxl.load_workbook(path, data_only=True, read_only=False)


def cell(v: Any) -> Any:
    """Normalize a cell to str | int | float | bool | None, blanks -> None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, dt.datetime):
        return v.isoformat()
    text = str(v).strip()
    if text == "" or text.lower() in {"n/a", "na", "----", "none", "null"}:
        return None
    return text


def sheet_rows(wb, name: str) -> list[dict[str, Any]]:
    """First row is the header; return a list of dicts keyed by header."""
    if wb is None or name not in wb.sheetnames:
        problems.append(f"missing sheet: {name}")
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [cell(h) for h in rows[0]]
    out = []
    for raw in rows[1:]:
        values = [cell(v) for v in raw]
        if all(v is None for v in values):
            continue
        out.append({h: values[i] if i < len(values) else None
                    for i, h in enumerate(header) if h})
    return out


def pick(row: dict, *keys: str) -> Any:
    for k in keys:
        if k in row and row[k] is not None:
            return row[k]
    return None


def as_int(v: Any) -> int | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"-?\d+", str(v))
    return int(m.group()) if m else None


def canonical_is(text: str | None) -> str:
    """
    Canonical key for an Indian Standard reference so that free text such as
    'IS 302 (Part 1)' can be resolved against 'IS 302 (Part 1):2024'.
    """
    if not text:
        return ""
    s = text.lower()
    s = re.sub(r"\s*/\s*iec[^ ]*\s*\d[\d\.\-\(\)\s]*", " ", s)   # drop IEC alias
    s = re.sub(r"\(\s*\d+(?:st|nd|rd|th)\s+revision\s*\)", " ", s)
    s = re.sub(r"\(\s*rev[^)]*\)", " ", s)
    s = re.sub(r":\d{4}", " ", s)
    s = re.sub(r"\b(19|20)\d{2}\b", " ", s)
    s = re.sub(r"[–——\-/,.;:]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def is_refs(text: str | None) -> list[str]:
    """All 'IS <n> [(part...)]' fragments inside a free-text cell."""
    if not text:
        return []
    return re.findall(r"IS\s?\d+(?:\s*\([^)]{0,60}\))?", text, flags=re.IGNORECASE)


def resolve_is(text: str | None, standards: list[dict]) -> tuple[list[str], list[str]]:
    """Return (resolved standard ids, unresolved IS fragments)."""
    resolved: list[str] = []
    unresolved: list[str] = []
    keys = [(s["id"], canonical_is(s.get("is_number"))) for s in standards]
    for frag in is_refs(text):
        frag_c = canonical_is(frag)
        if not frag_c:
            continue
        hits = [sid for sid, key in keys if key and (key == frag_c or key.startswith(frag_c + " "))]
        if not hits:
            # e.g. 'IS 8008 (Parts 1-9)' / 'IS 7634 (Part 1 & Part 2)':
            # fall back to the bare IS number and pull in every part.
            bare = re.sub(r"\(.*$", "", frag_c).strip()
            if bare and bare != frag_c:
                hits = [sid for sid, key in keys
                        if key and (key == bare or key.startswith(bare + " "))]
        if hits:
            for h in hits:
                if h not in resolved:
                    resolved.append(h)
        else:
            pretty = re.sub(r"\s+", " ", frag).strip()
            if pretty not in unresolved:
                unresolved.append(pretty)
    return resolved, unresolved


def source_category(source: dict) -> tuple[str, str]:
    """(category, authority) for the Sources & Provenance page."""
    sid = (source.get("id") or "").lower()
    url = (source.get("url") or "").lower()
    blob = " ".join(str(source.get(k) or "") for k in ("name", "raw_text", "what_it_gives")).lower()

    if sid.startswith("src-demo") or "demo" in blob and "synthetic" in blob:
        return "Demo", "Prototype demo fixture"
    if any(h in url for h in (
        "bis.gov.in", "services.bis.gov.in", "lims.bis.gov.in",
        "egazette.gov.in", "manakonline.in", "qco",
    )):
        return "Official BIS", "Official / Government"
    if sid.startswith("src-link") and url:
        return "Reference", "Third-party reference"
    if any(h in url for h in ("law.resource.org", "alephindia", "certification-india", "elplaw")):
        return "Reference", "Third-party reference"
    if any(h in blob for h in ("law.resource.org", "alephindia", "certification-india", "elplaw")):
        return "Reference", "Third-party reference"
    if "bis" in blob or "bis" in sid:
        # Named official BIS surface with no URL captured in the workbook.
        return "Official BIS", "Official / Government"
    return "Seed / Prototype", "Seed workbook note"


# --------------------------------------------------------------------------- #
# ingest
# --------------------------------------------------------------------------- #
def main() -> None:
    wb_primary = load(PRIMARY)
    wb_consumer = load(CONSUMER)
    wb_ref = load(REFERENCE)

    # ---- sources -------------------------------------------------------- #
    sources = []
    for row in sheet_rows(wb_primary, "Sources"):
        src = {
            "id": pick(row, "Source_ID"),
            "name": pick(row, "Source_Name"),
            "url": pick(row, "URL"),
            "status": pick(row, "Source_Status"),
            "verification_status": pick(row, "Verification_Status"),
            "what_it_gives": pick(row, "What_It_Gives"),
            "raw_text": pick(row, "Raw_Source_Text"),
            "workbook": PRIMARY,
            "sheet": "Sources",
        }
        cat, auth = source_category(src)
        src["category"] = cat
        src["authority"] = auth
        sources.append(src)
    source_ids = {s["id"] for s in sources}

    # The consumer demo records reference SRC-DEMO-CONS-001; if the workbook
    # does not register it, surface that honestly rather than inventing one.
    if "SRC-DEMO-CONS-001" not in source_ids:
        sources.append({
            "id": "SRC-DEMO-CONS-001",
            "name": "Prototype demo consumer record fixture",
            "url": None,
            "status": "Declared by Consumer_Demo_Records sheet",
            "verification_status": "Demo only — not verified against any BIS source",
            "what_it_gives": "Marks the three consumer demo records as synthetic UI fixtures.",
            "raw_text": "Referenced by Consumer_Demo_Records rows; no URL supplied in the workbook.",
            "category": "Demo",
            "authority": "Prototype demo fixture",
            "workbook": CONSUMER,
            "sheet": "Consumer_Demo_Records",
        })
        source_ids.add("SRC-DEMO-CONS-001")
        report.append("added missing demo source SRC-DEMO-CONS-001 (declared, no URL)")

    # A supplied reference workbook is itself a source — described, with no URL.
    sources.append({
        "id": "SRC-REF-STDLIST-001",
        "name": "BIS 'Total Published Standards' list (supplied workbook)",
        "url": None,
        "status": "Supplied with this project",
        "verification_status": "Supplied as-is; not re-verified in this prototype",
        "what_it_gives": "Broad catalogue of published Indian Standards used for exploration only — not used as evidence.",
        "raw_text": f"datasets/{REFERENCE}",
        "category": "Reference",
        "authority": "Reference corpus",
        "workbook": REFERENCE,
        "sheet": "Sheet1",
    })

    # ---- standards ------------------------------------------------------ #
    standards = []
    for row in sheet_rows(wb_primary, "Standards"):
        standards.append({
            "id": pick(row, "Standard_ID"),
            "domain": pick(row, "Domain"),
            "is_number": pick(row, "IS_Number"),
            "iec_reference": pick(row, "IEC_or_International_Reference"),
            "title": pick(row, "Title"),
            "standard_type": pick(row, "Standard_Type"),
            "status": pick(row, "Status"),
            "revision_or_year": pick(row, "Revision_or_Year"),
            "amendments_count": as_int(pick(row, "Amendments_Count")),
            "notes": pick(row, "Status_or_Notes_Raw"),
            "source_id": pick(row, "Source_ID"),
            "source_url": pick(row, "Source_URL"),
            "data_basis": pick(row, "Data_Basis"),
            "verification_status": pick(row, "Verification_Status"),
            "workbook": PRIMARY,
            "sheet": "Standards",
        })
    std_ids = {s["id"] for s in standards}

    # Amendment counts supplied by the legacy pipes sheet / Amendments sheet.
    amend_by_std: dict[str, int] = {}
    for row in sheet_rows(wb_primary, "Amendments"):
        sid = pick(row, "Standard_ID")
        n = as_int(pick(row, "Amendments_Count"))
        if sid and n is not None:
            amend_by_std[sid] = n
    for s in standards:
        if s["amendments_count"] is None and s["id"] in amend_by_std:
            s["amendments_count"] = amend_by_std[s["id"]]

    # ---- products ------------------------------------------------------- #
    legacy_sheets = {
        "Electrical Home Appliances": "Electrical_Appliances_Products",
        "Plastic/PVC Pipes & Fittings": "Plastic_Pipes_Products",
    }
    legacy_by_label: dict[str, dict] = {}
    for sheet in legacy_sheets.values():
        for row in sheet_rows(wb_primary, sheet):
            label = pick(row, "Product / Appliance Type", "Product / Use Case")
            if label:
                legacy_by_label[label.strip().lower()] = {
                    "applicable_standards_text": pick(row, "Applicable Standard(s)"),
                    "seed_notes": pick(row, "Notes for AI matching"),
                    "sheet": sheet,
                }

    products = []
    for row in sheet_rows(wb_primary, "Product_Categories"):
        label = pick(row, "Product_or_Use_Case") or ""
        legacy = legacy_by_label.get(label.strip().lower(), {})
        text = legacy.get("applicable_standards_text")
        resolved, unresolved = resolve_is(text, standards)
        products.append({
            "id": pick(row, "Product_ID"),
            "domain": pick(row, "Domain"),
            "name": label,
            "subcategory": pick(row, "Subcategory"),
            "material": pick(row, "Material"),
            "application": pick(row, "Application"),
            "source_notes": pick(row, "Source_Notes"),
            "derivation_note": pick(row, "Derivation_Note"),
            "applicable_standards_text": text,
            "seed_notes": legacy.get("seed_notes"),
            "seed_sheet": legacy.get("sheet"),
            "seed_standard_ids": resolved,
            "seed_unresolved_refs": unresolved,
            "workbook": PRIMARY,
            "sheet": "Product_Categories",
        })
    prod_ids = {p["id"] for p in products}

    matched = sum(1 for p in products if p["applicable_standards_text"])
    report.append(f"products joined to legacy applicable-standard text: {matched}/{len(products)}")

    # ---- relationships -------------------------------------------------- #
    relationships = []
    for row in sheet_rows(wb_primary, "Product_Standard_Relationships"):
        pid, sid = pick(row, "Product_ID"), pick(row, "Standard_ID")
        if pid not in prod_ids:
            problems.append(f"relationship {pick(row,'Relationship_ID')} -> unknown product {pid}")
        if sid not in std_ids:
            problems.append(f"relationship {pick(row,'Relationship_ID')} -> unknown standard {sid}")
        relationships.append({
            "id": pick(row, "Relationship_ID"),
            "product_id": pid,
            "standard_id": sid,
            "relationship_type": pick(row, "Relationship_Type"),
            "original_text": pick(row, "Original_Applicable_Standards_Text"),
            "source_notes": pick(row, "Source_Notes"),
            "derivation": pick(row, "Derivation"),
            "evidence_id": pick(row, "Evidence_ID"),
            "workbook": PRIMARY,
            "sheet": "Product_Standard_Relationships",
        })

    # ---- applicability rules -------------------------------------------- #
    rules = []
    for row in sheet_rows(wb_primary, "Applicability_Rules"):
        rules.append({
            "id": pick(row, "Rule_ID"),
            "product_id": pick(row, "Product_ID"),
            "standard_id": pick(row, "Standard_ID"),
            "field": pick(row, "Condition_Field"),
            "operator": pick(row, "Operator"),
            "value": pick(row, "Value"),
            "derivation": pick(row, "Derivation"),
            "validation_note": pick(row, "Validation_Note"),
            "evidence_id": pick(row, "Evidence_ID"),
            "workbook": PRIMARY,
            "sheet": "Applicability_Rules",
        })

    # ---- evidence ------------------------------------------------------- #
    evidence = []
    for row in sheet_rows(wb_primary, "Evidence"):
        evidence.append({
            "id": pick(row, "Evidence_ID"),
            "standard_id": pick(row, "Standard_ID"),
            "source_id": pick(row, "Source_ID"),
            "section": pick(row, "Section"),
            "clause": pick(row, "Clause"),
            "page": pick(row, "Page"),
            "text": pick(row, "Evidence_Text_or_Seed_Note"),
            "limitation": pick(row, "Evidence_Limitation"),
            "workbook": PRIMARY,
            "sheet": "Evidence",
        })

    # ---- supporting collections ---------------------------------------- #
    qcos = [{k: pick(r, k) for k in ("QCO_ID", "Category", "QCO_or_Certification_Status_Raw",
                                     "Key_Dates_Raw", "Source_ID", "Source_Raw", "Verification_Note")}
            for r in sheet_rows(wb_primary, "QCOs")]
    certifications = [{k: pick(r, k) for k in ("Certification_ID", "Standard_ID", "Scheme_Name",
                                               "Certification_Relevance", "Source_ID", "Note")}
                      for r in sheet_rows(wb_primary, "Certification_Schemes")]
    tests = [{k: pick(r, k) for k in ("Test_ID", "Standard_ID", "Test_or_Testing_Context",
                                      "Source_Note_Raw", "Source_ID", "Verification_Note")}
             for r in sheet_rows(wb_primary, "Testing_Requirements")]
    amendments = [{k: pick(r, k) for k in ("Amendment_Record_ID", "Standard_ID", "Amendments_Count",
                                           "Basis", "Amendment_Number_or_Date", "Note")}
                  for r in sheet_rows(wb_primary, "Amendments")]
    evaluation_cases = [{k: pick(r, k) for k in ("Case_ID", "Domain", "Input_Product_or_Query",
                                                 "Expected_or_Source_Mapped_Result",
                                                 "Why_This_Case_Matters", "Basis")}
                        for r in sheet_rows(wb_primary, "Evaluation_Cases")]
    source_links = [{"source": pick(r, "Source"), "url": pick(r, "URL"), "what_it_gives": pick(r, "What it gives you")}
                    for r in sheet_rows(wb_primary, "Source_Links")]
    qa_status = [{k: pick(r, k) for k in ("QA_ID", "Check", "Status", "Explanation")}
                 for r in sheet_rows(wb_primary, "QA_Status")]
    data_dictionary = [{"field": pick(r, "Field"), "meaning": pick(r, "Meaning"), "treatment": pick(r, "Treatment")}
                       for r in sheet_rows(wb_primary, "Data_Dictionary")]
    qco_notes = [{"category": pick(r, "Category"), "status": pick(r, "QCO / Certification status"),
                  "key_dates": pick(r, "Key dates"), "source": pick(r, "Source")}
                 for r in sheet_rows(wb_primary, "QCO_Notes")]

    # ---- consumer records ----------------------------------------------- #
    consumer_wb = wb_consumer if wb_consumer is not None else wb_primary
    consumer_workbook = CONSUMER if wb_consumer is not None else PRIMARY

    def consumer_rows(sheet: str) -> list[dict]:
        return sheet_rows(consumer_wb, sheet)

    consumer_records = []
    for row in consumer_rows("Consumer_Demo_Records") + consumer_rows("Consumer_BIS_Records"):
        is_demo = bool(row.get("Demo_Flag"))
        consumer_records.append({
            "id": pick(row, "Record_ID"),
            "identifier_type": pick(row, "Identifier_Type"),
            "identifier": pick(row, "Identifier"),
            "product": pick(row, "Product"),
            "manufacturer": pick(row, "Manufacturer"),
            "associated_is_number": pick(row, "Associated_IS_Number"),
            "standard_id": pick(row, "Standard_ID"),
            "certification_context": pick(row, "Certification_Context"),
            "status": pick(row, "Status"),
            "source_id": pick(row, "Source_ID"),
            "source_url": pick(row, "Source_URL"),
            "authority_level": pick(row, "Authority_Level"),
            "verification_status": pick(row, "Verification_Status"),
            "last_checked": pick(row, "Last_Checked"),
            "demo": is_demo,
            "notes": pick(row, "Notes"),
            "workbook": consumer_workbook,
            "sheet": "Consumer_Demo_Records" if is_demo else "Consumer_BIS_Records",
        })

    consumer_eval_cases = [{k: pick(r, k) for k in ("Case_ID", "Input_Type", "User_Input",
                                                    "Expected_Behavior", "Why_Test_Matters",
                                                    "Source_Basis", "Demo_or_Live")}
                           for r in consumer_rows("Consumer_Evaluation_Cases")]

    if not consumer_records:
        problems.append("no consumer records found")

    # ---- reference catalogue -------------------------------------------- #
    reference_standards: list[dict] = []
    if wb_ref is not None and "Sheet1" in wb_ref.sheetnames:
        ws = wb_ref["Sheet1"]
        for idx, raw in enumerate(ws.iter_rows(values_only=True)):
            if idx < 2:          # row 1 = title banner, row 2 = header
                continue
            vals = [cell(v) for v in raw]
            if len(vals) < 3 or not vals[1]:
                continue
            reference_standards.append({
                "sl_no": as_int(vals[0]),
                "is_number": vals[1],
                "title": vals[2],
                "aspect": vals[3] if len(vals) > 3 else None,
                "equivalence": vals[4] if len(vals) > 4 else None,
                "published": vals[5] if len(vals) > 5 else None,
                "action": vals[6] if len(vals) > 6 else None,
            })

    # ---- stats ---------------------------------------------------------- #
    domains = sorted({s["domain"] for s in standards if s["domain"]})
    stats = {
        "standards": len(standards),
        "products": len(products),
        "product_categories": len({p["name"] for p in products}),
        "domains": len(domains),
        "domain_list": domains,
        "relationships": len(relationships),
        "rules": len(rules),
        "sources": len(sources),
        "evidence": len(evidence),
        "consumer_records": len(consumer_records),
        "consumer_demo_records": sum(1 for c in consumer_records if c["demo"]),
        "qcos": len(qcos),
        "certifications": len(certifications),
        "testing_requirements": len(tests),
        "evaluation_cases": len(evaluation_cases) + len(consumer_eval_cases),
        "reference_standards": len(reference_standards),
    }

    kb = {
        "meta": {
            "name": "BIS Mitra",
            "tagline": "Indian Standards intelligence, made understandable.",
            "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
            "workbooks": [PRIMARY, CONSUMER, REFERENCE],
            "trust_policy": (
                "No standards, QCO dates, licence numbers, clauses, pages or source URLs "
                "were invented. Missing values remain null."
            ),
            "stats": stats,
        },
        "standards": standards,
        "products": products,
        "relationships": relationships,
        "rules": rules,
        "sources": sources,
        "evidence": evidence,
        "qcos": qcos,
        "qco_notes": qco_notes,
        "certifications": certifications,
        "testing_requirements": tests,
        "amendments": amendments,
        "evaluation_cases": evaluation_cases,
        "consumer_records": consumer_records,
        "consumer_evaluation_cases": consumer_eval_cases,
        "source_links": source_links,
        "reference_standards": reference_standards,
        "qa_status": qa_status,
        "data_dictionary": data_dictionary,
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(kb, ensure_ascii=False, indent=1), encoding="utf-8")

    print("=== ingest report ===")
    for line in report:
        print("  -", line)
    print("  - stats:", json.dumps(stats, ensure_ascii=False))
    print(f"  - wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size/1024:.1f} KB)")
    if problems:
        print("=== problems ===")
        for p in problems:
            print("  !", p)


if __name__ == "__main__":
    main()
